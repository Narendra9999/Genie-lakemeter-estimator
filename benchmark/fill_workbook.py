"""Write benchmark results into a copy of the DCP benchmark workbook (tab "C. Benchmark Q&A").

Usage: python3 fill_workbook.py <results.json> <input.xlsx> <output.xlsx>
Fills the empty "Agent answer", "Pass / Fail" and "Notes" columns and adds a sheet "D. Benchmark Run" with the detail.
"""
import json, sys
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

res_path, src, dst = sys.argv[1:4]
results = {r["n"]: r for r in json.load(open(res_path))}
wb = openpyxl.load_workbook(src)
ws = wb["C. Benchmark Q&A"]
hdr_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "#")
cols = {ws.cell(hdr_row, c).value: c for c in range(1, ws.max_column + 1)}
green, red = PatternFill("solid", fgColor="D9EAD3"), PatternFill("solid", fgColor="F4CCCC")
wrap = Alignment(wrap_text=True, vertical="top")
for r in range(hdr_row + 1, ws.max_row + 1):
    n = ws.cell(r, cols["#"]).value
    if not isinstance(n, int) or n not in results:
        continue
    x = results[n]
    answer = (x["agent_text"] or "").strip() or "(no text) result rows: " + json.dumps(x["agent_rows"][:4])[:600]
    ws.cell(r, cols["Agent answer"]).value = answer[:1500]
    ws.cell(r, cols["Pass / Fail"]).value = x["agent_pass"]
    ws.cell(r, cols["Pass / Fail"]).fill = green if x["agent_pass"] == "PASS" else red
    ws.cell(r, cols["Notes"]).value = (f"Agent vs calculator: {x['agent_notes']}. Calculator (FEVM rate card): {x['calculator_answer']}. "
                                       f"Engine at governed rates vs this workbook: {x['parity_v4']} {x['engine_governed_figures']}.")[:1500]
    for k in ("Agent answer", "Pass / Fail", "Notes"):
        ws.cell(r, cols[k]).alignment = wrap
d = wb.create_sheet("D. Benchmark Run")
heads = ["#", "Question", "Agent pass", "Agent notes", "Calculator answer (FEVM rate card)", "Calculator figures", "Engine figures at governed rates",
         "Workbook v4.0.2 figures", "Engine vs workbook", "Agent SQL", "Agent text", "Conversation id"]
d.append(heads)
for c in d[1]: c.font = Font(bold=True)
for n in sorted(results):
    x = results[n]
    d.append([n, x["question"], x["agent_pass"], x["agent_notes"], x["calculator_answer"], json.dumps(x["calculator_figures"]),
              json.dumps(x["engine_governed_figures"]), json.dumps(x["v4_expected_figures"]), x["parity_v4"],
              "\n".join(x["agent_sql"])[:2000], (x["agent_text"] or "")[:2000], x["conversation_id"]])
for col, w in zip("ABCDEFGHIJKL", (4, 40, 9, 40, 60, 22, 22, 22, 10, 60, 60, 36)):
    d.column_dimensions[col].width = w
for row in d.iter_rows(min_row=2):
    for c in row: c.alignment = wrap
wb.save(dst)
print("wrote", dst, "| agent PASS", sum(r["agent_pass"] == "PASS" for r in results.values()), "/", len(results),
      "| engine MATCH", sum(r["parity_v4"] == "MATCH" for r in results.values()), "/", len(results))
